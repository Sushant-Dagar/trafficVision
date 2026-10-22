"""Downloads of a finished video's results: an Excel workbook and CSV files.

Everything comes from the saved stats JSON, so exporting is instant.
"""
from __future__ import annotations

import csv
import io
import time
from typing import Any

TYPES = ["car", "motorcycle", "bus", "truck", "bicycle"]


def _road_name(stats: dict[str, Any], idx: Any) -> str:
    if idx is None:
        return ""
    return stats["roads"][idx]["name"]


def _not_counted_text(reason: str | None) -> str:
    return {
        "parked": "parked (never moved)",
        "brief": "seen under 0.5 s",
        "fragment": "picked up mid-junction",
        "outside": "never on a marked road",
    }.get(reason or "", reason or "")


def _exit_basis_text(basis: str | None) -> str:
    return {
        "seen": "seen on exit road",
        "linked": "seen on exit road (tracker re-found it)",
        "direction": "estimated from direction of travel",
    }.get(basis or "", "")


# ---------------------------------------------------------------------------
# Rows shared by the CSVs and the workbook
# ---------------------------------------------------------------------------
ROAD_HEADERS = [
    "Road", "Vehicles counted", "Cars", "Motorcycles", "Buses", "Trucks", "Bicycles",
    "2-wheelers", "4-wheelers", "Flow (veh/h)", "PCU", "PCU/h",
    "Left", "Straight", "Right", "U-turn", "Exit not seen",
    "Vehicles that stopped", "Avg wait (s)", "Max wait (s)",
    "Max queue (veh)", "Avg queue (veh)",
]


def road_rows(stats: dict[str, Any]) -> list[list[Any]]:
    rows = []
    for r in stats["roads"]:
        turns = r.get("turns") or {}
        queue = r.get("queue") or {}
        rows.append([
            r["name"], r["total"], *[r["by_type"].get(t, 0) for t in TYPES],
            r["categories"].get("2-wheeler", 0), r["categories"].get("4-wheeler", 0),
            r.get("flow_vph"), r.get("pcu"), r.get("pcu_ph"),
            turns.get("left", 0), turns.get("straight", 0), turns.get("right", 0), turns.get("U-turn", 0),
            r.get("not_seen_leaving"),
            r.get("stopped_count"), r.get("avg_wait"), r.get("max_wait"),
            queue.get("max"), queue.get("avg"),
        ])
    return rows


VEHICLE_HEADERS = [
    "Vehicle id", "Type", "Category", "Wheels", "Came from", "Left by", "Turn",
    "How exit was found", "Arrived (s)", "Waited (s)", "Wait cut off by clip",
    "PCU", "Detection confidence", "Not counted because",
]


def vehicle_rows(stats: dict[str, Any]) -> list[list[Any]]:
    rows = []
    for v in sorted(stats["vehicles"], key=lambda v: v["t"]):
        counted = v.get("road") is not None
        rows.append([
            v["id"], v["type"], v["category"], v["wheels"],
            _road_name(stats, v.get("road")), _road_name(stats, v.get("to")), v.get("turn") or "",
            _exit_basis_text(v.get("exit_basis")) if counted else "",
            v["t"], v.get("wait") if counted else "",
            ("yes" if v.get("wait_partial") else "no") if counted else "",
            v.get("pcu") if counted else "", round(v["confidence"], 3),
            _not_counted_text(v.get("not_counted")),
        ])
    return rows


def _csv(headers: list[str], rows: list[list[Any]]) -> bytes:
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow(headers)
    w.writerows(rows)
    # BOM so Excel opens it as UTF-8 (road names may contain any character).
    return ("﻿" + buf.getvalue()).encode("utf-8")


def roads_csv(stats: dict[str, Any]) -> bytes:
    return _csv(ROAD_HEADERS, road_rows(stats))


def vehicles_csv(stats: dict[str, Any]) -> bytes:
    return _csv(VEHICLE_HEADERS, vehicle_rows(stats))


# ---------------------------------------------------------------------------
# Excel workbook
# ---------------------------------------------------------------------------
def workbook(stats: dict[str, Any], video_name: str) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    bold = Font(bold=True)
    head_fill = PatternFill("solid", fgColor="F3B21F")
    title_font = Font(bold=True, size=14)
    note_font = Font(italic=True, color="666666")

    def header(ws, row: int, values: list[Any]) -> None:
        for col, val in enumerate(values, start=1):
            cell = ws.cell(row=row, column=col, value=val)
            cell.font = bold
            cell.fill = head_fill
            cell.alignment = Alignment(wrap_text=True, vertical="center")

    def fit(ws, min_w: int = 8, max_w: int = 40) -> None:
        for col in ws.columns:
            longest = max((len(str(c.value)) for c in col if c.value is not None), default=0)
            ws.column_dimensions[get_column_letter(col[0].column)].width = max(min_w, min(max_w, longest + 2))

    has_turns = "turning" in stats
    wb = Workbook()

    # --- Summary
    ws = wb.active
    ws.title = "Summary"
    ws["A1"] = f"Junction traffic survey - {video_name}"
    ws["A1"].font = title_font
    info = [
        ("Footage length (s)", stats.get("duration")),
        ("Resolution", stats.get("resolution")),
        ("Roads", len(stats["roads"])),
        ("Vehicles counted", stats.get("counted", stats.get("total"))),
        ("Flow, all roads (veh/h)", stats.get("flow_vph")),
        ("PCU, all roads", stats.get("pcu")),
        ("PCU/h, all roads", stats.get("pcu_ph")),
        ("Report generated", time.strftime("%Y-%m-%d %H:%M")),
    ]
    for i, (k, v) in enumerate(info, start=3):
        ws.cell(row=i, column=1, value=k).font = bold
        ws.cell(row=i, column=2, value=v)
    start = len(info) + 4
    header(ws, start, ROAD_HEADERS)
    for i, row in enumerate(road_rows(stats), start=start + 1):
        for j, val in enumerate(row, start=1):
            ws.cell(row=i, column=j, value=val)
    note_row = start + len(stats["roads"]) + 2
    ws.cell(row=note_row, column=1, value=(
        "Flow and PCU/h scale this clip's counts to one hour - for a short clip they are "
        "rough. Each vehicle is counted once, on the road it came from."
    )).font = note_font
    fit(ws)
    ws.column_dimensions["A"].width = 26
    ws.freeze_panes = ws.cell(row=start + 1, column=2)

    # --- Turning movements
    if has_turns:
        tm = stats["turning"]
        names = [r["name"] for r in stats["roads"]]
        ws = wb.create_sheet("Turning movements")
        ws["A1"] = "Turning movements (vehicles)"
        ws["A1"].font = title_font
        header(ws, 3, ["From \\ To", *names, "Exit not seen", "Total"])
        for i, name in enumerate(names):
            r = 4 + i
            ws.cell(row=r, column=1, value=name).font = bold
            for j in range(len(names)):
                ws.cell(row=r, column=2 + j, value=tm["matrix"][i][j] if i != j or tm["matrix"][i][j] else "")
            ws.cell(row=r, column=2 + len(names), value=tm["not_seen_leaving"][i])
            ws.cell(row=r, column=3 + len(names), value=sum(tm["matrix"][i]) + tm["not_seen_leaving"][i])

        r0 = 6 + len(names)
        ws.cell(row=r0, column=1, value="Movement type").font = title_font
        header(ws, r0 + 1, ["From \\ To", *names])
        for i, name in enumerate(names):
            ws.cell(row=r0 + 2 + i, column=1, value=name).font = bold
            for j in range(len(names)):
                ws.cell(row=r0 + 2 + i, column=2 + j, value=tm["labels"][i][j])

        r1 = r0 + 3 + len(names)
        n_est = sum(map(sum, tm.get("estimated", [])))
        ws.cell(row=r1, column=1, value=(
            f"{n_est} of these movements were estimated from the vehicle's direction of travel "
            "(it disappeared, e.g. behind a building, before reaching its exit road). "
            "Left/right are from the driver's point of view."
        )).font = note_font
        fit(ws, max_w=24)
        ws.column_dimensions["A"].width = 26

    # --- Vehicles
    ws = wb.create_sheet("Vehicles")
    header(ws, 1, VEHICLE_HEADERS)
    for i, row in enumerate(vehicle_rows(stats), start=2):
        for j, val in enumerate(row, start=1):
            ws.cell(row=i, column=j, value=val)
    fit(ws)
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = ws.dimensions

    # --- Queue over time
    if has_turns and stats["roads"] and stats["roads"][0].get("queue"):
        ws = wb.create_sheet("Queue over time")
        step = stats["roads"][0]["queue"]["step"]
        header(ws, 1, ["Time (s)", *[f"{r['name']} - stopped vehicles" for r in stats["roads"]]])
        n = max(len(r["queue"]["values"]) for r in stats["roads"])
        for k in range(n):
            ws.cell(row=2 + k, column=1, value=round(k * step, 2))
            for j, r in enumerate(stats["roads"]):
                vals = r["queue"]["values"]
                ws.cell(row=2 + k, column=2 + j, value=vals[k] if k < len(vals) else None)
        fit(ws, max_w=30)
        ws.freeze_panes = "B2"

    # --- Method
    ws = wb.create_sheet("Method")
    ws["A1"] = "How these numbers are worked out"
    ws["A1"].font = title_font
    lines = [
        "Detection: YOLOv8 (COCO classes car, motorcycle, bus, truck, bicycle); tracking: ByteTrack.",
        "Came from: the road outline a vehicle's ground point (bottom-centre of its box) is inside "
        "within ~1 s of first being seen.",
        "Not counted: parked vehicles, detections under 0.5 s, vehicles picked up mid-junction "
        "(tracker re-found one, would double count), vehicles never on a marked road.",
        "Left by: the last other road it was seen on; if the tracker lost it, the id that picked "
        "it up; if it disappeared mid-junction, the road it was heading for (estimate).",
        "Stopped: moving slower than 0.2 of its own width per second. Wait = time stopped on the "
        "road it came from. Vehicles that stopped for under 1 s aren't in the average wait.",
        "Queue: vehicles stopped on the road at that moment.",
        "PCU: IRC:106-1990 (urban roads), interpolated by each type's share of counted traffic "
        "(lower value at 5%, higher at 10% or more).",
    ]
    for i, line in enumerate(lines, start=3):
        ws.cell(row=i, column=1, value=line)
    if stats.get("pcu_factors"):
        r = len(lines) + 4
        header(ws, r, ["Vehicle type", "PCU used"])
        for i, (t, f) in enumerate(stats["pcu_factors"].items(), start=r + 1):
            ws.cell(row=i, column=1, value=t)
            ws.cell(row=i, column=2, value=f)
    ws.column_dimensions["A"].width = 110

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
